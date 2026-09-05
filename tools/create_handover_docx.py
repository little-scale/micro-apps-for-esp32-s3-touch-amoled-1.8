#!/usr/bin/env python3
"""Build the polished developer handover DOCX from DEVELOPER_HANDOVER.md."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "DEVELOPER_HANDOVER.md"
OUTPUT_DIR = ROOT / "docs"
OUTPUT = OUTPUT_DIR / "ESP32-S3-Classroom-Controller-Developer-Handover.docx"
ASSET_DIR = OUTPUT_DIR / "handover-assets"

BLACK = RGBColor(0x12, 0x17, 0x1C)
MID = RGBColor(0x43, 0x52, 0x5E)
ACCENT = RGBColor(0x00, 0x71, 0x89)
LIGHT = "E9F1F3"
PALE = "F4F7F8"
GRID = "B9C6CC"


def font_path(bold: bool = False) -> str:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return ""


def pil_font(size: int, bold: bool = False):
    path = font_path(bold)
    return ImageFont.truetype(path, size) if path else ImageFont.load_default()


def rounded_box(draw, xy, fill, outline=(18, 23, 28), width=4, radius=18):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, fill=(0, 113, 137), width=5):
    draw.line([start, end], fill=fill, width=width)
    x2, y2 = end
    x1, y1 = start
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    head = 16
    wing = 8
    points = [
        (x2, y2),
        (x2 - ux * head + px * wing, y2 - uy * head + py * wing),
        (x2 - ux * head - px * wing, y2 - uy * head - py * wing),
    ]
    draw.polygon(points, fill=fill)


def centered_text(draw, box, text, font, fill=(18, 23, 28), spacing=7):
    left, top, right, bottom = box
    lines = text.split("\n")
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = top + (bottom - top - total) / 2
    for line, height in zip(lines, heights):
        bounds = draw.textbbox((0, 0), line, font=font)
        width = bounds[2] - bounds[0]
        draw.text((left + (right - left - width) / 2, y), line, font=font, fill=fill)
        y += height + spacing


def create_screen_diagram(path: Path):
    image = Image.new("RGB", (1200, 760), "white")
    draw = ImageDraw.Draw(image)
    title = pil_font(38, True)
    label = pil_font(27, True)
    small = pil_font(22)
    tiny = pil_font(18)
    draw.text((55, 34), "Stable 368 × 448 screen framework", font=title,
              fill=(18, 23, 28))

    x, y, w, h = 90, 115, 368, 448
    rounded_box(draw, (x, y, x + w, y + h), fill=(9, 12, 15), outline=(18, 23, 28),
                width=9, radius=56)
    draw.rectangle((x + 18, y + 10, x + w - 18, y + 43), fill=(28, 35, 40))
    for offset, color in [(0, (0, 176, 205)), (42, (255, 255, 255)),
                          (84, (255, 72, 176)), (280, (255, 255, 255))]:
        draw.rounded_rectangle((x + 30 + offset, y + 17, x + 58 + offset, y + 38),
                               radius=4, fill=color)

    cx0, cy0, cw, ch = x + 20, y + 57, 328, 310
    draw.rectangle((cx0, cy0, cx0 + cw, cy0 + ch), fill=(255, 255, 255))
    draw.rectangle((cx0 + 150, cy0, cx0 + 177, cy0 + ch), fill=(9, 12, 15))
    draw.rectangle((cx0, cy0 + 141, cx0 + cw, cy0 + 168), fill=(9, 12, 15))
    draw.rectangle((cx0 + 150, cy0 + 141, cx0 + 177, cy0 + 168), fill=(255, 255, 255))
    px0, py0, pw, ph = x + 20, y + 385, 328, 49
    draw.rectangle((px0, py0, px0 + pw, py0 + ph), fill=(245, 245, 245))
    for index in range(8):
        fill = (0, 113, 137) if index == 0 else (90, 100, 106)
        xx = px0 + 64 + index * 25
        draw.rectangle((xx, py0 + 18, xx + 11, py0 + 29), fill=fill)

    draw.text((535, 130), "Status strip", font=label, fill=(18, 23, 28))
    draw.text((535, 172), "43 px high", font=small, fill=(67, 82, 94))
    arrow(draw, (515, 170), (455, 150))
    draw.text((535, 276), "Performance canvas", font=label, fill=(18, 23, 28))
    draw.text((535, 318), "328 × 310 px", font=small, fill=(67, 82, 94))
    draw.text((535, 352), "PSRAM composited and aligned", font=small, fill=(67, 82, 94))
    arrow(draw, (515, 340), (455, 350))
    draw.text((535, 485), "Next page control", font=label, fill=(18, 23, 28))
    draw.text((535, 527), "328 × 49 px visible", font=small, fill=(67, 82, 94))
    draw.text((535, 560), "expanded lower-edge touch target", font=small, fill=(67, 82, 94))
    arrow(draw, (515, 515), (455, 520))

    draw.line((535, 632, 1120, 632), fill=(0, 113, 137), width=4)
    draw.text((535, 654), "Current order", font=label, fill=(18, 23, 28))
    draw.text((535, 700), "Keyboard  •  Buttons  •  XY  •  Faders  •  Balls  •\n"
              "Pendulums  •  Particles  •  Spectrum", font=tiny, fill=(67, 82, 94),
              spacing=8)
    image.save(path, quality=95)


def create_architecture_diagram(path: Path):
    image = Image.new("RGB", (1600, 880), "white")
    draw = ImageDraw.Draw(image)
    title = pil_font(43, True)
    box_title = pil_font(30, True)
    box_body = pil_font(22)
    draw.text((60, 35), "One event router keeps interaction and transport separate",
              font=title, fill=(18, 23, 28))

    sources = [
        ((70, 150, 390, 275), "Touch", "CST820 gesture lifecycle"),
        ((70, 330, 390, 455), "Motion", "QMI8658 screen frame"),
        ((70, 510, 390, 635), "Sound", "ES8311 energy and FFT"),
        ((70, 690, 390, 815), "Remote input", "OSC or BLE validation"),
    ]
    for box, heading, body in sources:
        rounded_box(draw, box, fill=(244, 247, 248), outline=(185, 198, 204), width=4)
        centered_text(draw, (box[0], box[1] + 4, box[2], box[1] + 58), heading,
                      box_title)
        centered_text(draw, (box[0], box[1] + 58, box[2], box[3] - 4), body, box_body,
                      fill=(67, 82, 94))
        arrow(draw, (box[2], (box[1] + box[3]) // 2), (565, 440))

    central = (565, 285, 1015, 595)
    rounded_box(draw, central, fill=(223, 240, 243), outline=(0, 113, 137), width=6,
                radius=26)
    centered_text(draw, (565, 308, 1015, 414), "UserInterface and\nshared ControlState",
                  box_title)
    centered_text(draw, (595, 430, 985, 555),
                  "Updates state\nQueues local events\nDraws only changed regions",
                  box_body, fill=(67, 82, 94))

    outputs = [
        ((1210, 205, 1530, 350), "Display", "Composited regions"),
        ((1210, 415, 1530, 560), "OSC", "UDP message output"),
        ((1210, 625, 1530, 770), "BLE", "GATT notifications"),
    ]
    for box, heading, body in outputs:
        rounded_box(draw, box, fill=(244, 247, 248), outline=(185, 198, 204), width=4)
        centered_text(draw, (box[0], box[1] + 8, box[2], box[1] + 70), heading,
                      box_title)
        centered_text(draw, (box[0], box[1] + 70, box[2], box[3] - 6), body, box_body,
                      fill=(67, 82, 94))
        arrow(draw, (1015, 440), (box[0], (box[1] + box[3]) // 2))

    draw.text((1040, 362), "Local events only", font=box_body, fill=(0, 113, 137))
    draw.text((825, 802), "Remote values stop at state and display  •  no echo",
              font=pil_font(25, True), fill=(18, 23, 28))
    image.save(path, quality=95)


def create_no_echo_diagram(path: Path):
    image = Image.new("RGB", (1500, 520), "white")
    draw = ImageDraw.Draw(image)
    title = pil_font(42, True)
    heading = pil_font(29, True)
    body = pil_font(22)
    draw.text((55, 35), "No echo is an architectural boundary", font=title,
              fill=(18, 23, 28))
    boxes = [
        (55, 150, 330, 300, "Local touch", "owns the gesture"),
        (485, 150, 795, 300, "Shared state", "updates the display"),
        (960, 100, 1425, 225, "OSC and BLE output", "local changes only"),
        (960, 320, 1425, 445, "Remote OSC or BLE", "validated input only"),
    ]
    for left, top, right, bottom, head, sub in boxes:
        rounded_box(draw, (left, top, right, bottom), fill=(244, 247, 248),
                    outline=(185, 198, 204), width=4)
        centered_text(draw, (left, top + 10, right, top + 75), head, heading)
        centered_text(draw, (left, top + 72, right, bottom - 5), sub, body,
                      fill=(67, 82, 94))
    arrow(draw, (330, 225), (485, 225))
    arrow(draw, (795, 205), (960, 163))
    arrow(draw, (960, 382), (795, 255))
    draw.line((835, 300, 900, 365), fill=(197, 54, 54), width=8)
    draw.line((900, 300, 835, 365), fill=(197, 54, 54), width=8)
    draw.text((805, 410), "never forward remote input", font=body,
              fill=(197, 54, 54))
    image.save(path, quality=95)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "007189")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(
    r"(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)"
)


def add_inline(paragraph, text: str, base_bold: bool = False):
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.bold = base_bold
        token = match.group(0)
        if token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                label, url = link.groups()
                if url.startswith("http://") or url.startswith("https://"):
                    add_hyperlink(paragraph, label, url)
                else:
                    run = paragraph.add_run(label)
                    run.bold = base_bold
                    run.font.color.rgb = ACCENT
            else:
                paragraph.add_run(token)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Menlo"
            run.font.size = Pt(8.6)
            run.font.color.rgb = RGBColor(0x16, 0x42, 0x51)
            run.bold = base_bold
        else:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = base_bold


def add_bottom_rule(paragraph, color="007189", size="18"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def set_repeat_and_keep(paragraph, keep_next=False, keep_lines=True):
    fmt = paragraph.paragraph_format
    fmt.keep_with_next = keep_next
    fmt.keep_together = keep_lines


def style_document(document: Document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, before, after in [
        ("Title", 28, 0, 12),
        ("Heading 1", 20, 16, 7),
        ("Heading 2", 14.5, 13, 5),
        ("Heading 3", 11.2, 9, 3),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Menlo"
    code.font.size = Pt(7.3)
    code.font.color.rgb = RGBColor(0x13, 0x2E, 0x38)
    code.paragraph_format.left_indent = Inches(0.17)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(9.2)
        style.paragraph_format.space_after = Pt(2.5)


def add_header_footer(document: Document):
    for section in document.sections:
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("CLASSROOM CONTROLLER   ·   DEVELOPER HANDOVER")
        run.font.name = "Arial"
        run.font.size = Pt(7.2)
        run.font.bold = True
        run.font.color.rgb = MID
        add_bottom_rule(paragraph, color="B9C6CC", size="6")

        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("5 SEPTEMBER 2026   ·   ")
        run.font.name = "Arial"
        run.font.size = Pt(7.2)
        run.font.color.rgb = MID
        add_page_number(paragraph)


def add_cover(document: Document, screen_path: Path):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(32)
    run = paragraph.add_run("ESP32 S3 TOUCH AMOLED")
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = ACCENT

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(9)
    title.add_run("Classroom Controller\nDeveloper Handover")
    add_bottom_rule(title, color="007189", size="24")

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run(
        "A verified foundation for touch, motion, microphone, OSC over Wi-Fi, "
        "Bluetooth Low Energy, Max and Ableton-oriented classroom projects"
    )
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = MID

    meta = document.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    labels = [
        ("TARGET", "Waveshare ESP32-S3-Touch-AMOLED-1.8 V2"),
        ("DISPLAY", "368 × 448 portrait AMOLED  ·  CO5300  ·  CST820"),
        ("BASELINE", "Eight pages  ·  keyboard first  ·  radar retired"),
        ("BUILD", "1,384,495 bytes  ·  44 percent of application partition"),
    ]
    for row, (left, right) in zip(meta.rows, labels):
        row.cells[0].width = Inches(1.12)
        row.cells[1].width = Inches(5.85)
        set_cell_shading(row.cells[0], "007189")
        set_cell_shading(row.cells[1], "F4F7F8")
        set_cell_margins(row.cells[0], 80, 100, 80, 100)
        set_cell_margins(row.cells[1], 80, 100, 80, 100)
        p = row.cells[0].paragraphs[0]
        r = p.add_run(left)
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        p = row.cells[1].paragraphs[0]
        r = p.add_run(right)
        r.font.name = "Arial"
        r.font.size = Pt(8.6)
        r.font.bold = True
        r.font.color.rgb = BLACK

    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(14)
    picture.add_run().add_picture(str(screen_path), width=Inches(6.9))

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Project baseline  ·  5 September 2026")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = MID
    document.add_page_break()


def add_guide_map(document: Document):
    heading = document.add_paragraph("Guide map", style="Heading 1")
    add_bottom_rule(heading, color="007189", size="16")
    intro = document.add_paragraph(
        "This handover is designed to be useful at three levels: a quick restart for a "
        "future development session, a hardware and protocol reference, and a worked guide "
        "for adding a student page."
    )
    intro.paragraph_format.space_after = Pt(9)
    rows = [
        ("ORIENT", "Current status  ·  product intent  ·  hardware version warning"),
        ("BUILD", "Repository map  ·  build and flash  ·  recovery"),
        ("INTERFACE", "Page order  ·  visual language  ·  touch  ·  redraw strategy"),
        ("SENSORS", "IMU frame  ·  movement trigger  ·  microphone  ·  FFT"),
        ("CONNECT", "Settings  ·  identity  ·  OSC  ·  BLE  ·  no echo"),
        ("EXTEND", "Retired radar example  ·  new-page recipe  ·  verification"),
        ("HAND OFF", "Known limits  ·  next steps  ·  reusable context block"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.18)
        cells[1].width = Inches(5.75)
        set_cell_shading(cells[0], "E9F1F3")
        set_cell_shading(cells[1], "F8FAFA")
        for cell in cells:
            set_cell_margins(cell, 100, 110, 100, 110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[0].paragraphs[0]
        r = p.add_run(left)
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        p = cells[1].paragraphs[0]
        r = p.add_run(right)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = BLACK
    document.add_paragraph()


def add_code_block(document: Document, lines: list[str]):
    paragraph = document.add_paragraph(style="Code Block")
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "007189")
    borders.append(left)
    p_pr.append(borders)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Menlo"
    run.font.size = Pt(7.3)
    paragraph.paragraph_format.keep_together = len(lines) <= 14


def add_table(document: Document, rows: list[list[str]]):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            repeat_table_header(row)
        for column_index in range(columns):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = values[column_index] if column_index < len(values) else ""
            if row_index == 0:
                set_cell_shading(cell, "173540")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F4F7F8")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, text, base_bold=row_index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(7.9 if columns >= 3 else 8.3)
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [part.strip() for part in stripped.split("|")]


def is_table_separator(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def render_markdown(document: Document, markdown: str, architecture_path: Path,
                    no_echo_path: Path):
    lines = markdown.splitlines()
    index = 0
    paragraph_buffer: list[str] = []
    first_heading_skipped = False

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph = document.add_paragraph()
            add_inline(paragraph, " ".join(line.strip() for line in paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            if level == 1 and not first_heading_skipped:
                first_heading_skipped = True
                index += 1
                continue
            if text == "Quick context block for a future development session":
                document.add_page_break()
            style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
            paragraph = document.add_paragraph(text, style=style)
            if level == 2:
                add_bottom_rule(paragraph, color="007189", size="10")
            if text == "Repository map":
                pic = document.add_paragraph()
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(str(architecture_path), width=Inches(7.0))
                cap = document.add_paragraph("Figure 1  Software and data flow")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].font.size = Pt(7.8)
                cap.runs[0].italic = True
                cap.runs[0].font.color.rgb = MID
            if text == "No echo and ownership rule":
                pic = document.add_paragraph()
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(str(no_echo_path), width=Inches(7.0))
                cap = document.add_paragraph("Figure 2  Remote state terminates at the interface")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].font.size = Pt(7.8)
                cap.runs[0].italic = True
                cap.runs[0].font.color.rgb = MID
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            flush_paragraph()
            rows = [parse_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if bullet or numbered:
            flush_paragraph()
            if bullet:
                paragraph = document.add_paragraph(style="List Bullet")
                add_inline(paragraph, bullet.group(1))
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.23)
                paragraph.paragraph_format.first_line_indent = Inches(-0.23)
                paragraph.paragraph_format.space_after = Pt(2.5)
                number_run = paragraph.add_run(numbered.group(1) + ".  ")
                number_run.bold = True
                number_run.font.color.rgb = ACCENT
                add_inline(paragraph, numbered.group(2))
            index += 1
            continue

        if stripped == "":
            flush_paragraph()
            index += 1
            continue

        # The baseline facts already appear as a designed table on the cover.
        if stripped.startswith("**Project baseline:") or stripped.startswith("**Target hardware:") \
                or stripped.startswith("**Display:") or stripped.startswith("**Touch:") \
                or stripped.startswith("**Current firmware:") \
                or stripped.startswith("**Verified build size:") \
                or stripped.startswith("**Verified RAM use:"):
            index += 1
            continue
        paragraph_buffer.append(stripped.rstrip("  "))
        index += 1

    flush_paragraph()


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)
    screen_path = ASSET_DIR / "screen-framework.png"
    architecture_path = ASSET_DIR / "architecture.png"
    no_echo_path = ASSET_DIR / "no-echo.png"
    create_screen_diagram(screen_path)
    create_architecture_diagram(architecture_path)
    create_no_echo_diagram(no_echo_path)

    document = Document()
    style_document(document)
    add_header_footer(document)
    add_cover(document, screen_path)
    add_guide_map(document)
    render_markdown(document, SOURCE.read_text(encoding="utf-8"), architecture_path,
                    no_echo_path)

    properties = document.core_properties
    properties.title = "ESP32 S3 Touch AMOLED Classroom Controller Developer Handover"
    properties.subject = "Firmware architecture, protocols, build workflow and extension guide"
    properties.author = "Classroom Controller project"
    properties.keywords = "ESP32-S3, Waveshare, AMOLED, OSC, BLE, IMU, Max, Ableton"
    properties.comments = "Generated from DEVELOPER_HANDOVER.md"

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
